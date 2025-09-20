import os
import io
import uuid
import json
import csv
import html
from html.parser import HTMLParser
from pathlib import Path
from typing import List, Dict, Tuple, Optional

from utils.helpers.chroma_db import ChromaVectorStore


def _read_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def _read_pdf(path: Path) -> str:
    try:
        import pdfplumber
    except ImportError:
        raise RuntimeError("pdfplumber is required for PDF support. Install with: pip install pdfplumber")
    text_parts: List[str] = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            txt = page.extract_text() or ""
            if txt:
                text_parts.append(txt)
    return "\n".join(text_parts)


def _read_docx(path: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError:
        raise RuntimeError("python-docx is required for DOCX support. Install with: pip install python-docx")
    doc = docx.Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text)



def _read_xlsx(path: Path) -> str:
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError("openpyxl is required for XLSX support. Install with: pip install openpyxl")
    wb = openpyxl.load_workbook(str(path), data_only=True)
    parts: List[str] = []
    for ws in wb.worksheets:
        parts.append(f"# Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            line = ", ".join(cells).strip()
            if line:
                parts.append(line)
    return "\n".join(parts)


def _read_csv(path: Path) -> str:
    with path.open("r", encoding="utf-8", errors="ignore", newline="") as f:
        reader = csv.reader(f)
        rows = []
        for row in reader:
            rows.append(", ".join(row))
    return "\n".join(rows)


class _HTMLTextExtractor(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts: List[str] = []

    def handle_data(self, data: str):
        data = data.strip()
        if data:
            self.parts.append(html.unescape(data))

    def text(self) -> str:
        return "\n".join(self.parts)


def _read_html(path: Path) -> str:
    content = path.read_text(encoding="utf-8", errors="ignore")
    parser = _HTMLTextExtractor()
    parser.feed(content)
    return parser.text()


def _read_json(path: Path) -> str:
    data = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
    try:
        return json.dumps(data, ensure_ascii=False, indent=2)
    except Exception:
        return str(data)


def _extract_text_any(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in [".txt", ".md", ".log", ".rst"]:
        return _read_txt(path)
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".docx":
        return _read_docx(path)
    if suffix == ".xlsx":
        return _read_xlsx(path)
    if suffix == ".csv":
        return _read_csv(path)
    if suffix in [".html", ".htm"]:
        return _read_html(path)
    if suffix == ".json":
        return _read_json(path)
    # Fallback: try as text
    return _read_txt(path)


def chunk_by_word_count(text: str, chunk_size: int = 500, overlap: int = 100) -> List[str]:
    if chunk_size <= 0:
        raise ValueError("chunk_size must be > 0")
    if overlap < 0 or overlap >= chunk_size:
        overlap = 0
    words = text.split()
    n = len(words)
    chunks: List[str] = []
    start = 0
    while start < n:
        end = min(start + chunk_size, n)
        chunk_words = words[start:end]
        chunk = " ".join(chunk_words).strip()
        if chunk:
            chunks.append(chunk)
        if end == n:
            break
        start = end - overlap
    return chunks


def add_documents(document_path: Path, chunk_size: int = 500):
    if not document_path.exists():
        raise FileNotFoundError(f"Document not found: {document_path}")

    # Extract text from many common formats
    document_data = _extract_text_any(document_path).strip()
    if not document_data:
        return {"added": 0, "message": "No extractable text found in file"}

    # Chunk by word count for better semantic retrieval windows
    chunks = chunk_by_word_count(document_data, chunk_size=chunk_size, overlap=100)
    if not chunks:
        return {"added": 0, "message": "No text chunks produced"}

    # Prepare metadata and IDs
    store = ChromaVectorStore()
    base_id = f"{document_path.stem}-{uuid.uuid4().hex[:8]}"
    ids = [f"{base_id}-{i}" for i in range(len(chunks))]
    metadatas: List[Dict] = [
        {
            "source": document_path.name,
            "chunk_index": i,
            "total_chunks": len(chunks),
            "path": str(document_path.resolve()),
            "ext": document_path.suffix.lower(),
        }
        for i in range(len(chunks))
    ]

    ok = store.add_documents(chunks, metadatas=metadatas, ids=ids)
    return {
        "ok": ok,
        "added": len(chunks) if ok else 0,
        "collection_info": store.info(),
    }
